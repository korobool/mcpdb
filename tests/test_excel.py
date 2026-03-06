"""Tests for v2.2.x — Excel Support.

Covers CellGrid, WorkbookReader, structure analyzer, ExcelSource,
and Excel MCP tools (load_excel, list_sheets, read_range,
extract_excel_table, explain_formula).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from mcpdb.sources.excel.cell_grid import CellGrid, CellValue, MergedRegion, _classify_cell_type
from mcpdb.sources.excel.structure_analyzer import (
    DetectedTable,
    TextRegion,
    analyze_sheet,
    _detect_tables,
    _detect_text_regions,
    _find_header_row,
)
from mcpdb.sources.excel.workbook_reader import WorkbookReader
from mcpdb.sources.excel.source import ExcelSource
from mcpdb.sources.base import SourceType
from mcpdb.sources.registry import SourceRegistry
from mcpdb.mcp.tools.excel import (
    get_excel_tools,
    handle_excel_tool,
    _basic_formula_explanation,
)


# ---------------------------------------------------------------------------
# Fixtures — Excel file creation
# ---------------------------------------------------------------------------

@pytest.fixture
def simple_workbook(tmp_path) -> Path:
    """Create a simple workbook with one data table."""
    path = tmp_path / "simple.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Region", "Amount"])
    ws.append(["Widget A", "North", 100.50])
    ws.append(["Widget B", "South", 200.75])
    ws.append(["Gadget X", "East", 75.25])
    ws.append(["Widget A", "West", 150.00])
    wb.save(path)
    return path


@pytest.fixture
def multi_sheet_workbook(tmp_path) -> Path:
    """Create a workbook with multiple sheets."""
    path = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Products table
    ws1 = wb.active
    ws1.title = "Products"
    ws1.append(["ID", "Name", "Price"])
    ws1.append([1, "Widget A", 10.00])
    ws1.append([2, "Widget B", 20.00])
    ws1.append([3, "Gadget X", 15.00])

    # Sheet 2: Orders table
    ws2 = wb.create_sheet("Orders")
    ws2.append(["OrderID", "ProductID", "Qty", "Total"])
    ws2.append([101, 1, 5, 50.00])
    ws2.append([102, 2, 3, 60.00])
    ws2.append([103, 1, 10, 100.00])

    wb.save(path)
    return path


@pytest.fixture
def mixed_workbook(tmp_path) -> Path:
    """Create a workbook with text + table (mixed structure)."""
    path = tmp_path / "mixed.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    # Text region (rows 1-2)
    ws.append(["Q4 2024 Sales Report"])
    ws.append(["Generated on 2024-12-15"])

    # Empty row (row 3)
    ws.append([])

    # Table region (rows 4-7)
    ws.append(["Region", "Revenue", "Customers"])
    ws.append(["North", 50000, 120])
    ws.append(["South", 35000, 85])
    ws.append(["East", 42000, 95])

    wb.save(path)
    return path


@pytest.fixture
def formula_workbook(tmp_path) -> Path:
    """Create a workbook with formulas."""
    path = tmp_path / "formulas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Q1", "Q2", "Total"])
    ws.append(["Salaries", 50000, 52000, None])
    ws.append(["Marketing", 10000, 12000, None])
    ws.append(["Total", None, None, None])

    # Add formulas
    ws["D2"] = "=B2+C2"
    ws["D3"] = "=B3+C3"
    ws["B4"] = "=SUM(B2:B3)"
    ws["C4"] = "=SUM(C2:C3)"
    ws["D4"] = "=SUM(D2:D3)"

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# CellValue tests
# ---------------------------------------------------------------------------

class TestCellValue:
    def test_empty_display(self):
        cv = CellValue()
        assert cv.display == ""

    def test_text_display(self):
        cv = CellValue(value="hello", dtype="text")
        assert cv.display == "hello"

    def test_number_display(self):
        cv = CellValue(value=42.5, dtype="number")
        assert cv.display == "42.5"

    def test_none_display(self):
        cv = CellValue(value=None)
        assert cv.display == ""


class TestClassifyCellType:
    def test_none(self):
        assert _classify_cell_type(None) == "empty"

    def test_bool(self):
        assert _classify_cell_type(True) == "bool"
        assert _classify_cell_type(False) == "bool"

    def test_number(self):
        assert _classify_cell_type(42) == "number"
        assert _classify_cell_type(3.14) == "number"

    def test_text(self):
        assert _classify_cell_type("hello") == "text"

    def test_formula(self):
        assert _classify_cell_type("=SUM(A1:A10)") == "formula"

    def test_other(self):
        from datetime import datetime
        assert _classify_cell_type(datetime.now()) == "text"


# ---------------------------------------------------------------------------
# CellGrid tests
# ---------------------------------------------------------------------------

class TestCellGrid:
    def test_empty_grid(self):
        grid = CellGrid()
        assert grid.row_count == 0
        assert grid.col_count == 0
        assert grid.to_text() == "(empty)"

    def test_get_valid(self):
        cells = [
            [CellValue(value="A"), CellValue(value="B")],
            [CellValue(value=1), CellValue(value=2)],
        ]
        grid = CellGrid(cells=cells, row_count=2, col_count=2)
        assert grid.get(0, 0).value == "A"
        assert grid.get(1, 1).value == 2

    def test_get_out_of_bounds(self):
        grid = CellGrid(cells=[[CellValue(value="X")]], row_count=1, col_count=1)
        assert grid.get(5, 5).value is None

    def test_to_text(self):
        cells = [
            [CellValue(value="Name"), CellValue(value="Age")],
            [CellValue(value="Alice"), CellValue(value=30)],
        ]
        grid = CellGrid(cells=cells, row_count=2, col_count=2)
        text = grid.to_text()
        assert "Name" in text
        assert "Alice" in text
        assert "30" in text

    def test_to_text_truncation(self):
        long_val = "x" * 50
        cells = [[CellValue(value=long_val)]]
        grid = CellGrid(cells=cells, row_count=1, col_count=1)
        text = grid.to_text()
        assert "..." in text
        assert len(text) < len(long_val)

    def test_to_text_max_rows(self):
        cells = [[CellValue(value=f"row{i}")] for i in range(100)]
        grid = CellGrid(cells=cells, row_count=100, col_count=1)
        text = grid.to_text(max_rows=10)
        assert "90 more rows" in text


class TestMergedRegion:
    def test_size(self):
        mr = MergedRegion(min_row=0, min_col=0, max_row=2, max_col=3)
        assert mr.size == (3, 4)


# ---------------------------------------------------------------------------
# Structure analyzer tests
# ---------------------------------------------------------------------------

class TestFindHeaderRow:
    def _make_grid(self, rows: list[list]) -> CellGrid:
        cells = []
        for row in rows:
            cells.append([CellValue(value=v, dtype=_classify_cell_type(v)) for v in row])
        return CellGrid(
            cells=cells,
            row_count=len(cells),
            col_count=len(cells[0]) if cells else 0,
        )

    def test_valid_header(self):
        grid = self._make_grid([
            ["Name", "Age", "City"],
            ["Alice", 30, "NYC"],
        ])
        result = _find_header_row(grid, 0)
        assert result is not None
        start, end, headers = result
        assert headers == ["Name", "Age", "City"]

    def test_single_cell_not_header(self):
        grid = self._make_grid([
            ["Title", None, None],
            ["Data", 1, 2],
        ])
        result = _find_header_row(grid, 0)
        assert result is None

    def test_gap_too_large(self):
        grid = self._make_grid([
            ["A", None, None, None, "B"],
            [1, None, None, None, 2],
        ])
        result = _find_header_row(grid, 0)
        assert result is None

    def test_no_data_below(self):
        grid = self._make_grid([
            ["Name", "Age"],
            [None, None],
        ])
        result = _find_header_row(grid, 0)
        assert result is None


class TestDetectTables:
    def _make_grid(self, rows: list[list]) -> CellGrid:
        cells = []
        max_cols = max(len(r) for r in rows) if rows else 0
        for row in rows:
            padded = row + [None] * (max_cols - len(row))
            cells.append([CellValue(value=v, dtype=_classify_cell_type(v)) for v in padded])
        return CellGrid(
            cells=cells,
            row_count=len(cells),
            col_count=max_cols,
        )

    def test_simple_table(self):
        grid = self._make_grid([
            ["Name", "Age", "City"],
            ["Alice", 30, "NYC"],
            ["Bob", 25, "LA"],
        ])
        tables = _detect_tables(grid, "Sheet1")
        assert len(tables) == 1
        assert tables[0].headers == ["Name", "Age", "City"]
        assert tables[0].row_count == 2

    def test_no_table(self):
        grid = self._make_grid([
            ["Just some text"],
            ["And more text"],
        ])
        tables = _detect_tables(grid, "Sheet1")
        assert len(tables) == 0

    def test_empty_grid(self):
        grid = CellGrid()
        result = analyze_sheet(grid, "Empty")
        assert result["structure_type"] == "empty"
        assert result["tables"] == []
        assert result["text_regions"] == []


class TestAnalyzeSheet:
    def _make_grid(self, rows: list[list]) -> CellGrid:
        cells = []
        max_cols = max(len(r) for r in rows) if rows else 0
        for row in rows:
            padded = row + [None] * (max_cols - len(row))
            cells.append([CellValue(value=v, dtype=_classify_cell_type(v)) for v in padded])
        return CellGrid(
            cells=cells,
            row_count=len(cells),
            col_count=max_cols,
        )

    def test_table_type(self):
        grid = self._make_grid([
            ["A", "B"],
            [1, 2],
            [3, 4],
        ])
        result = analyze_sheet(grid, "Sheet1")
        assert result["structure_type"] == "table"
        assert len(result["tables"]) == 1

    def test_text_type(self):
        grid = self._make_grid([
            ["Title of the document"],
            ["Line two of text"],
        ])
        result = analyze_sheet(grid, "Sheet1")
        assert result["structure_type"] == "text"
        assert len(result["text_regions"]) > 0


class TestDetectedTable:
    def test_range_str(self):
        t = DetectedTable(
            sheet_name="Sheet1",
            header_row=0,
            data_start_row=1,
            data_end_row=3,
            start_col=0,
            end_col=2,
            headers=["A", "B", "C"],
            row_count=3,
        )
        r = t.range_str
        assert r == "A1:C4"


# ---------------------------------------------------------------------------
# WorkbookReader tests
# ---------------------------------------------------------------------------

class TestWorkbookReader:
    def test_sheet_names(self, simple_workbook):
        reader = WorkbookReader(simple_workbook)
        assert reader.sheet_names == ["Sales"]
        reader.close()

    def test_list_sheets(self, simple_workbook):
        reader = WorkbookReader(simple_workbook)
        sheets = reader.list_sheets()
        assert len(sheets) == 1
        assert sheets[0]["name"] == "Sales"
        assert sheets[0]["max_row"] >= 5
        reader.close()

    def test_read_sheet(self, simple_workbook):
        reader = WorkbookReader(simple_workbook)
        grid = reader.read_sheet("Sales")
        assert grid.row_count >= 5
        assert grid.col_count >= 3
        assert grid.get(0, 0).value == "Product"
        reader.close()

    def test_read_sheet_not_found(self, simple_workbook):
        reader = WorkbookReader(simple_workbook)
        with pytest.raises(ValueError, match="not found"):
            reader.read_sheet("NonExistent")
        reader.close()

    def test_read_range(self, simple_workbook):
        reader = WorkbookReader(simple_workbook)
        grid = reader.read_range("Sales", "A1:C3")
        assert grid.row_count == 3
        assert grid.col_count == 3
        reader.close()

    def test_multi_sheet(self, multi_sheet_workbook):
        reader = WorkbookReader(multi_sheet_workbook)
        assert len(reader.sheet_names) == 2
        assert "Products" in reader.sheet_names
        assert "Orders" in reader.sheet_names
        reader.close()

    def test_get_formulas(self, formula_workbook):
        reader = WorkbookReader(formula_workbook)
        formulas = reader.get_formulas("Budget")
        assert len(formulas) >= 5
        cells = {f["cell"] for f in formulas}
        assert "D2" in cells
        assert "B4" in cells
        reader.close()


# ---------------------------------------------------------------------------
# ExcelSource tests
# ---------------------------------------------------------------------------

class TestExcelSource:
    def test_source_type(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        assert source.source_type == SourceType.EXCEL
        source.close()

    def test_metadata(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        meta = source.get_metadata()
        assert meta.source_type == SourceType.EXCEL
        assert "simple" in meta.source_id
        assert len(meta.tables) >= 1
        assert "sql" in meta.capabilities
        source.close()

    def test_context_description(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        desc = source.get_context_description()
        assert "Excel Workbook" in desc
        assert "Sales" in desc
        source.close()

    def test_list_sheets(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        sheets = source.list_sheets()
        assert len(sheets) == 1
        assert sheets[0]["name"] == "Sales"
        assert "structure_type" in sheets[0]
        source.close()

    def test_read_range(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        text = source.read_range("Sales")
        assert "Product" in text
        assert "Widget A" in text
        source.close()

    def test_get_sheet_structure(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        structure = source.get_sheet_structure("Sales")
        assert "Sales" in structure
        source.close()

    def test_get_sheet_structure_not_found(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        with pytest.raises(ValueError, match="not found"):
            source.get_sheet_structure("NonExistent")
        source.close()

    def test_extract_table(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        df = source.extract_table("Sales", "A1:C5")
        assert isinstance(df, pd.DataFrame)
        assert len(df) >= 3
        source.close()

    def test_multi_sheet_tables(self, multi_sheet_workbook):
        source = ExcelSource(multi_sheet_workbook)
        meta = source.get_metadata()
        # Should detect tables in both sheets
        assert len(meta.tables) >= 2
        source.close()

    def test_mixed_structure(self, mixed_workbook):
        source = ExcelSource(mixed_workbook)
        sheets = source.list_sheets()
        assert sheets[0].get("structure_type") in ("table", "mixed", "text")
        source.close()

    def test_close_idempotent(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        source.close()
        source.close()  # Should not raise

    def test_display_name(self, simple_workbook):
        source = ExcelSource(simple_workbook, display_name="My Report")
        assert source._display_name == "My Report"
        source.close()

    def test_get_formulas(self, formula_workbook):
        source = ExcelSource(formula_workbook)
        formulas = source.get_formulas("Budget")
        assert len(formulas) >= 5
        source.close()


# ---------------------------------------------------------------------------
# SQL querying of extracted tables
# ---------------------------------------------------------------------------

class TestExcelSQL:
    def test_query_extracted_table(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        meta = source.get_metadata()
        assert len(meta.tables) >= 1

        table_name = meta.tables[0].name
        df = source._backend.execute_query(f"SELECT * FROM {table_name}")
        assert len(df) >= 4
        source.close()

    def test_sql_aggregation(self, simple_workbook):
        source = ExcelSource(simple_workbook)
        meta = source.get_metadata()
        table_name = meta.tables[0].name

        df = source._backend.execute_query(
            f"SELECT COUNT(*) as cnt FROM {table_name}"
        )
        assert df.iloc[0]["cnt"] >= 4
        source.close()

    def test_multi_sheet_sql(self, multi_sheet_workbook):
        source = ExcelSource(multi_sheet_workbook)
        meta = source.get_metadata()
        assert len(meta.tables) >= 2

        for t in meta.tables:
            df = source._backend.execute_query(f"SELECT COUNT(*) as cnt FROM {t.name}")
            assert df.iloc[0]["cnt"] >= 3
        source.close()


# ---------------------------------------------------------------------------
# Excel MCP tools tests
# ---------------------------------------------------------------------------

class TestExcelTools:
    def test_tool_definitions(self):
        tools = get_excel_tools()
        names = {t.name for t in tools}
        assert "load_excel" in names
        assert "list_sheets" in names
        assert "read_range" in names
        assert "extract_excel_table" in names
        assert "explain_formula" in names
        assert len(tools) == 5

    async def test_load_excel(self, simple_workbook):
        registry = SourceRegistry()
        result = await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook)},
            registry,
        )
        assert len(result) == 1
        text = result[0].text
        assert "Loaded Excel" in text
        assert "Source ID" in text
        assert len(registry.list_sources()) == 1

    async def test_load_excel_with_display_name(self, simple_workbook):
        registry = SourceRegistry()
        result = await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook), "display_name": "My Sales"},
            registry,
        )
        assert "My Sales" in result[0].text

    async def test_load_excel_bad_path(self):
        registry = SourceRegistry()
        result = await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": "/nonexistent/file.xlsx"},
            registry,
        )
        assert "Error" in result[0].text

    async def test_list_sheets(self, simple_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "list_sheets",
            {"_name": "list_sheets", "source_id": source_id},
            registry,
        )
        assert "Sales" in result[0].text

    async def test_list_sheets_not_found(self):
        registry = SourceRegistry()
        result = await handle_excel_tool(
            "list_sheets",
            {"_name": "list_sheets", "source_id": "nonexistent"},
            registry,
        )
        assert "not found" in result[0].text

    async def test_read_range_full_sheet(self, simple_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "read_range",
            {"_name": "read_range", "source_id": source_id, "sheet": "Sales"},
            registry,
        )
        assert "Product" in result[0].text
        assert "Widget" in result[0].text

    async def test_read_range_specific(self, simple_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "read_range",
            {"_name": "read_range", "source_id": source_id, "sheet": "Sales", "range": "A1:B3"},
            registry,
        )
        text = result[0].text
        assert "Product" in text

    async def test_extract_excel_table(self, simple_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(simple_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "extract_excel_table",
            {"_name": "extract_excel_table", "source_id": source_id, "sheet": "Sales", "range": "A1:C5"},
            registry,
        )
        text = result[0].text
        assert "Extracted table" in text
        assert "rows" in text

    async def test_explain_formula_list_all(self, formula_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(formula_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "explain_formula",
            {"_name": "explain_formula", "source_id": source_id, "sheet": "Budget"},
            registry,
        )
        text = result[0].text
        assert "Formulas" in text
        assert "SUM" in text or "D2" in text

    async def test_explain_formula_specific_cell(self, formula_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(formula_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "explain_formula",
            {"_name": "explain_formula", "source_id": source_id, "sheet": "Budget", "cell": "B4"},
            registry,
        )
        text = result[0].text
        assert "B4" in text
        assert "SUM" in text

    async def test_explain_formula_no_formula(self, formula_workbook):
        registry = SourceRegistry()
        await handle_excel_tool(
            "load_excel",
            {"_name": "load_excel", "file_path": str(formula_workbook)},
            registry,
        )
        source_id = registry.list_sources()[0].source_id

        result = await handle_excel_tool(
            "explain_formula",
            {"_name": "explain_formula", "source_id": source_id, "sheet": "Budget", "cell": "A1"},
            registry,
        )
        assert "No formula" in result[0].text

    async def test_unknown_tool(self):
        registry = SourceRegistry()
        result = await handle_excel_tool("unknown_tool", {"_name": "unknown_tool"}, registry)
        assert "Unknown" in result[0].text


# ---------------------------------------------------------------------------
# Basic formula explanation tests
# ---------------------------------------------------------------------------

class TestBasicFormulaExplanation:
    def test_sum(self):
        result = _basic_formula_explanation("=SUM(A1:A10)")
        assert "Sums" in result or "Sum" in result

    def test_average(self):
        result = _basic_formula_explanation("=AVERAGE(B1:B20)")
        assert "average" in result.lower()

    def test_count(self):
        result = _basic_formula_explanation("=COUNT(C1:C50)")
        assert "count" in result.lower() or "Count" in result

    def test_if(self):
        result = _basic_formula_explanation("=IF(A1>10,\"yes\",\"no\")")
        assert "Conditional" in result or "if" in result.lower()

    def test_vlookup(self):
        result = _basic_formula_explanation("=VLOOKUP(A1,B1:C10,2,FALSE)")
        assert "look" in result.lower()

    def test_unknown_formula(self):
        result = _basic_formula_explanation("=CUSTOMFUNC(X1)")
        assert "CUSTOMFUNC" in result

    def test_max(self):
        result = _basic_formula_explanation("=MAX(D1:D100)")
        assert "maximum" in result.lower() or "max" in result.lower()

    def test_min(self):
        result = _basic_formula_explanation("=MIN(E1:E50)")
        assert "minimum" in result.lower() or "min" in result.lower()


# ---------------------------------------------------------------------------
# Integration: ExcelSource registered in SourceRegistry
# ---------------------------------------------------------------------------

class TestExcelRegistration:
    def test_register_and_query(self, simple_workbook):
        registry = SourceRegistry()
        source = ExcelSource(simple_workbook)
        source_id = registry.add_source(source)

        retrieved = registry.get_source(source_id)
        assert retrieved.source_type == SourceType.EXCEL

        meta = retrieved.get_metadata()
        assert meta.source_type == SourceType.EXCEL
        assert len(meta.tables) >= 1
        source.close()

    def test_remove_excel_source(self, simple_workbook):
        registry = SourceRegistry()
        source = ExcelSource(simple_workbook)
        source_id = registry.add_source(source)
        registry.remove_source(source_id)
        assert len(registry.list_sources()) == 0

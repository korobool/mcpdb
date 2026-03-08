"""Tests for the MCPDB web client (FastAPI backend)."""

from __future__ import annotations

import json
import pytest
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient


@pytest.fixture
def app():
    """Create a test app with mocked LLM and no database."""
    with patch("mcpdb.llm.factory.create_provider") as mock_create:
        mock_create.return_value = MagicMock()
        from mcpdb.web.app import create_app
        yield create_app()


@pytest.fixture
def client(app):
    return TestClient(app)


class TestHealth:
    def test_health_endpoint(self, client):
        resp = client.get("/api/health")
        assert resp.status_code == 200
        data = resp.json()
        assert "version" in data
        assert "source_count" in data

    def test_index_page(self, client):
        resp = client.get("/")
        assert resp.status_code == 200
        assert "MCPDB" in resp.text
        assert "WebSocket" in resp.text

    def test_no_auto_connect_from_env(self, client):
        """MCPDB_DATABASE_URL should NOT auto-connect, only prefill."""
        resp = client.get("/api/sources")
        assert resp.json() == []  # nothing auto-connected

    def test_favicon(self, client):
        resp = client.get("/favicon.ico")
        assert resp.status_code == 200
        assert "svg" in resp.text


class TestSources:
    def test_list_sources_empty(self, client):
        resp = client.get("/api/sources")
        assert resp.status_code == 200
        assert resp.json() == []

    def test_connect_database_invalid(self, client):
        resp = client.post("/api/sources/database", json={
            "connection_string": "invalid://nope",
        })
        assert resp.status_code == 400

    def test_connect_database_sqlite(self, client, tmp_path):
        db_path = tmp_path / "test.db"
        import sqlite3
        conn = sqlite3.connect(str(db_path))
        conn.execute("CREATE TABLE test (id INTEGER PRIMARY KEY, name TEXT)")
        conn.execute("INSERT INTO test VALUES (1, 'hello')")
        conn.commit()
        conn.close()

        resp = client.post("/api/sources/database", json={
            "connection_string": f"sqlite:///{db_path}",
            "display_name": "testdb",
        })
        assert resp.status_code == 200
        data = resp.json()
        assert "source_id" in data
        assert "metadata" in data

        # Verify it shows up in list
        resp = client.get("/api/sources")
        sources = resp.json()
        assert len(sources) == 1
        assert sources[0]["display_name"] == "testdb"
        assert "table_summary" in sources[0]

    def test_remove_source_not_found(self, client):
        resp = client.delete("/api/sources/nonexistent")
        assert resp.status_code == 404


class TestERD:
    def test_erd_with_source(self, client, tmp_path):
        db_path = tmp_path / "erd_test.db"
        import sqlite3
        conn = sqlite3.connect(str(db_path))
        conn.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT)")
        conn.execute("CREATE TABLE orders (id INTEGER PRIMARY KEY, user_id INTEGER)")
        conn.commit()
        conn.close()

        client.post("/api/sources/database", json={
            "connection_string": f"sqlite:///{db_path}",
        })

        sources = client.get("/api/sources").json()
        sid = sources[0]["source_id"]

        resp = client.get(f"/api/sources/{sid}/erd")
        assert resp.status_code == 200
        data = resp.json()
        assert "mermaid" in data
        assert "erDiagram" in data["mermaid"]


class TestWebSocket:
    def test_ws_empty_question(self, client):
        with client.websocket_connect("/ws/query") as ws:
            ws.send_json({"type": "query", "question": ""})
            resp = ws.receive_json()
            assert resp["type"] == "error"
            assert "Empty" in resp["message"]

    def test_ws_no_sources(self, client):
        with client.websocket_connect("/ws/query") as ws:
            ws.send_json({"type": "query", "question": "test"})
            resp = ws.receive_json()
            assert resp["type"] == "error"
            assert "sources" in resp["message"].lower()

    def test_ws_unknown_type(self, client):
        with client.websocket_connect("/ws/query") as ws:
            ws.send_json({"type": "garbage"})
            resp = ws.receive_json()
            assert resp["type"] == "error"

    def test_ws_invalid_json(self, client):
        with client.websocket_connect("/ws/query") as ws:
            ws.send_text("not json at all")
            resp = ws.receive_json()
            assert resp["type"] == "error"
            assert "JSON" in resp["message"]

    def test_ws_cancel_no_query(self, client):
        with client.websocket_connect("/ws/query") as ws:
            ws.send_json({"type": "cancel"})
            resp = ws.receive_json()
            assert resp["type"] == "error"
            assert "No active" in resp["message"]


class TestCSVUpload:
    def test_upload_single_csv(self, client, tmp_path):
        csv_path = tmp_path / "sales.csv"
        csv_path.write_text("id,product,amount\n1,Widget,100\n2,Gadget,200\n")

        with open(csv_path, "rb") as f:
            resp = client.post(
                "/api/sources/csv",
                files=[("files", ("sales.csv", f, "text/csv"))],
                data={"display_name": "test-csv"},
            )
        assert resp.status_code == 200
        data = resp.json()
        assert "source_id" in data
        assert "metadata" in data

        # Verify it shows up in list
        sources = client.get("/api/sources").json()
        assert len(sources) == 1
        assert sources[0]["display_name"] == "test-csv"
        assert sources[0]["source_type"] == "csv"

    def test_upload_multiple_csvs(self, client, tmp_path):
        csv1 = tmp_path / "orders.csv"
        csv1.write_text("order_id,customer,total\n1,Alice,50\n2,Bob,75\n")
        csv2 = tmp_path / "customers.csv"
        csv2.write_text("name,city\nAlice,NYC\nBob,LA\n")

        files = []
        handles = []
        for p in [csv1, csv2]:
            h = open(p, "rb")
            handles.append(h)
            files.append(("files", (p.name, h, "text/csv")))
        try:
            resp = client.post("/api/sources/csv", files=files)
        finally:
            for h in handles:
                h.close()

        assert resp.status_code == 200
        data = resp.json()
        assert "source_id" in data

        # Should have tables for both files
        sources = client.get("/api/sources").json()
        assert len(sources) == 1
        tables = sources[0].get("tables", [])
        assert len(tables) == 2

    def test_upload_csv_no_files(self, client):
        resp = client.post("/api/sources/csv")
        assert resp.status_code == 422  # validation error — no files

    def test_upload_csv_auto_name(self, client, tmp_path):
        csv_path = tmp_path / "data.csv"
        csv_path.write_text("a,b\n1,2\n")

        with open(csv_path, "rb") as f:
            resp = client.post(
                "/api/sources/csv",
                files=[("files", ("data.csv", f, "text/csv"))],
            )
        assert resp.status_code == 200
        sources = client.get("/api/sources").json()
        assert "data.csv" in sources[0]["display_name"]

    def test_remove_csv_source(self, client, tmp_path):
        csv_path = tmp_path / "remove_me.csv"
        csv_path.write_text("x,y\n1,2\n")

        with open(csv_path, "rb") as f:
            resp = client.post(
                "/api/sources/csv",
                files=[("files", ("remove_me.csv", f, "text/csv"))],
            )
        assert resp.status_code == 200
        sid = resp.json()["source_id"]

        resp = client.delete(f"/api/sources/{sid}")
        assert resp.status_code == 200

        sources = client.get("/api/sources").json()
        assert len(sources) == 0

    def test_csv_sample_endpoint(self, client, tmp_path):
        csv_path = tmp_path / "sample_test.csv"
        csv_path.write_text("id,name,value\n1,alpha,10\n2,beta,20\n3,gamma,30\n")

        with open(csv_path, "rb") as f:
            resp = client.post(
                "/api/sources/csv",
                files=[("files", ("sample_test.csv", f, "text/csv"))],
            )
        assert resp.status_code == 200
        sid = resp.json()["source_id"]

        # Get table name from sources list
        sources = client.get("/api/sources").json()
        table_name = sources[0]["tables"][0]["name"]

        resp = client.get(f"/api/sources/{sid}/sample/{table_name}?limit=2")
        assert resp.status_code == 200
        data = resp.json()
        assert "columns" in data
        assert "rows" in data
        assert len(data["rows"]) == 2


class TestExcelUpload:
    def test_upload_excel(self, client, tmp_path):
        import openpyxl
        path = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Value"])
        ws.append(["Alpha", 10])
        ws.append(["Beta", 20])
        wb.save(path)

        with open(path, "rb") as f:
            resp = client.post(
                "/api/sources/excel",
                files=[("file", ("test.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
                data={"display_name": "test-excel"},
            )
        assert resp.status_code == 200
        data = resp.json()
        assert "source_id" in data
        assert "metadata" in data

        sources = client.get("/api/sources").json()
        assert any(s["source_type"] == "excel" for s in sources)

    def test_excel_is_queryable(self, client, tmp_path):
        import openpyxl
        path = tmp_path / "queryable.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Id", "Score"])
        ws.append([1, 90])
        ws.append([2, 85])
        wb.save(path)

        with open(path, "rb") as f:
            resp = client.post(
                "/api/sources/excel",
                files=[("file", ("queryable.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
            )
        assert resp.status_code == 200
        sid = resp.json()["source_id"]

        # Should be able to get sample from extracted table
        sources = client.get("/api/sources").json()
        excel_source = [s for s in sources if s["source_id"] == sid][0]
        tables = excel_source.get("tables", [])
        if tables:
            tname = tables[0]["name"]
            sample = client.get(f"/api/sources/{sid}/sample/{tname}?limit=5")
            assert sample.status_code == 200


class TestExport:
    """Test export endpoints."""

    SAMPLE_TURNS = [{
        "question": "How many tables?",
        "response": "There are **21 tables** in the database.",
        "artifacts": [],
        "executionLog": [
            {"step_number": 1, "action": "sql",
             "code": "SELECT COUNT(*) FROM information_schema.tables",
             "status": "success", "result_summary": "21 tables"},
        ],
        "mode": "quick",
        "elapsed": 2.5,
    }]

    def test_export_html(self, client):
        resp = client.post("/api/export/html", json={"turns": self.SAMPLE_TURNS, "dbInfo": "bgins"})
        assert resp.status_code == 200
        data = resp.json()
        assert "export_id" in data
        assert "url" in data
        # Fetch the actual HTML via the view endpoint
        view_resp = client.get(data["url"])
        assert view_resp.status_code == 200
        assert "text/html" in view_resp.headers["content-type"]
        assert "How many tables" in view_resp.text
        assert "21 tables" in view_resp.text

    def test_export_html_empty_turns(self, client):
        resp = client.post("/api/export/html", json={"turns": []})
        assert resp.status_code == 400

    def test_export_notebook(self, client):
        resp = client.post("/api/export/notebook", json={
            "turns": self.SAMPLE_TURNS,
            "connectionString": "mysql+pymysql://test@localhost/db",
            "dbType": "mysql",
            "dbInfo": "bgins",
        })
        assert resp.status_code == 200
        assert "ipynb" in resp.headers.get("content-disposition", "")
        # Verify it's valid notebook JSON
        nb = json.loads(resp.content)
        assert "cells" in nb
        assert len(nb["cells"]) >= 3  # title + setup + at least 1 turn

    def test_export_notebook_empty_turns(self, client):
        resp = client.post("/api/export/notebook", json={"turns": []})
        assert resp.status_code == 400

    def test_export_html_with_artifacts(self, client):
        turns = [{
            "question": "Show chart",
            "response": "Here is the chart.",
            "artifacts": [{
                "type": "chart", "title": "Test",
                "data": {"image_base64": "iVBORw0KGgo=", "format": "png"},
            }, {
                "type": "table", "title": "Data",
                "data": {"columns": ["A", "B"], "rows": [[1, 2]]},
            }],
            "executionLog": [],
            "mode": "simple",
            "elapsed": 3.0,
        }]
        resp = client.post("/api/export/html", json={"turns": turns})
        assert resp.status_code == 200
        data = resp.json()
        view_resp = client.get(data["url"])
        assert view_resp.status_code == 200
        assert "data:image/png;base64" in view_resp.text
        assert "<table" in view_resp.text

    def test_export_view_one_time_use(self, client):
        """Export view endpoint should only serve the HTML once."""
        resp = client.post("/api/export/html", json={"turns": self.SAMPLE_TURNS})
        url = resp.json()["url"]
        # First access works
        assert client.get(url).status_code == 200
        # Second access returns 404
        assert client.get(url).status_code == 404


class TestLLMConfig:
    def test_get_config(self, client):
        resp = client.get("/api/config/llm")
        assert resp.status_code == 200
        data = resp.json()
        assert "provider" in data
        assert "model" in data

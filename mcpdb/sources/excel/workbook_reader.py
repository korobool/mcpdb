"""Workbook reader — openpyxl wrapper for reading Excel files."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import openpyxl

from mcpdb.sources.excel.cell_grid import CellGrid

logger = logging.getLogger("mcpdb.excel.reader")


class WorkbookReader:
    """Reads Excel workbooks using openpyxl."""

    def __init__(self, file_path: str | Path) -> None:
        self._path = str(file_path)

        # Load in data-only mode (computed values)
        self._wb_data = openpyxl.load_workbook(self._path, data_only=True, read_only=True)
        # Load in formula mode
        try:
            self._wb_formula = openpyxl.load_workbook(self._path, data_only=False, read_only=True)
        except Exception:
            self._wb_formula = None

        logger.info("Loaded workbook: %s (%d sheets)", self._path, len(self._wb_data.sheetnames))

    @property
    def sheet_names(self) -> list[str]:
        return self._wb_data.sheetnames

    def list_sheets(self) -> list[dict[str, Any]]:
        """Return metadata for all sheets."""
        result = []
        for name in self.sheet_names:
            ws = self._wb_data[name]
            result.append({
                "name": name,
                "max_row": ws.max_row or 0,
                "max_column": ws.max_column or 0,
            })
        return result

    def read_sheet(self, sheet_name: str) -> CellGrid:
        """Read a sheet as a CellGrid (data values)."""
        if sheet_name not in self._wb_data.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found")
        ws = self._wb_data[sheet_name]
        return CellGrid.from_worksheet(ws, data_only=True)

    def read_range(
        self,
        sheet_name: str,
        range_str: str | None = None,
        include_formulas: bool = False,
    ) -> CellGrid:
        """Read a specific range from a sheet."""
        if sheet_name not in self._wb_data.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found")

        wb = self._wb_formula if (include_formulas and self._wb_formula) else self._wb_data
        ws = wb[sheet_name]

        if range_str is None:
            return CellGrid.from_worksheet(ws)

        # Parse range like "A1:D10"
        from openpyxl.utils import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(range_str)

        from mcpdb.sources.excel.cell_grid import CellValue, _classify_cell_type
        cells = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                                 min_col=min_col, max_col=max_col):
            row_cells = []
            for cell in row:
                cv = CellValue(
                    value=cell.value,
                    dtype=_classify_cell_type(cell.value),
                )
                row_cells.append(cv)
            cells.append(row_cells)

        return CellGrid(
            cells=cells,
            row_count=len(cells),
            col_count=len(cells[0]) if cells else 0,
        )

    def get_formulas(self, sheet_name: str) -> list[dict[str, str]]:
        """List all formulas in a sheet."""
        if self._wb_formula is None:
            return []

        ws = self._wb_formula[sheet_name]
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                    })
        return formulas

    def close(self) -> None:
        """Close the workbook."""
        try:
            self._wb_data.close()
        except Exception:
            pass
        if self._wb_formula:
            try:
                self._wb_formula.close()
            except Exception:
                pass
